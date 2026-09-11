"""
Pruebas del lector de la planilla de asistencia (matriz por bloques).

Generan un .xlsx con la estructura EXACTA del reporte del reloj:
    fila 0: titulo | fila 2: periodo en la columna C | fila 3: dias 1..31
    fila 4, 6, 8...: 'ID:' (col 0) + ID (col 2) + 'Nombre:' (col 8) + nombre (col 10)
                     + 'Departamento:' (col 18) + valor (col 20)
    fila 5, 7, 9...: marcaciones alineadas a los dias

Ejecutar:  python -m pytest scripts/parse_asistencia_test.py -q
"""

import datetime as dt
import json

import openpyxl
import pytest

from parse_asistencia import (
    a_json,
    a_largo,
    a_matriz,
    detectar_cabecera_dias,
    detectar_periodo,
    extraer_horas,
    leer_grilla,
    main,
    parsear_reporte,
)

PERIODO = "2026-08-01 ~ 2026-08-31"


def construir_reporte(ruta, empleados, offset_dias=4, con_periodo=True, cabecera_repetida=False):
    """
    Escribe un reporte con el formato del reloj.

    `offset_dias` es la columna (1-based) donde arranca el dia 1: con 1 los dias
    empiezan en la columna 0 (indice 0) y con 5 en la columna 4.
    """
    libro = openpyxl.Workbook()
    hoja = libro.active
    hoja.title = "Eventos"

    hoja.cell(row=1, column=1, value="Reporte de Eventos de Asistencia")
    if con_periodo:
        hoja.cell(row=3, column=3, value=PERIODO)
    for dia in range(1, 32):
        hoja.cell(row=4, column=offset_dias + dia - 1, value=dia)

    fila = 5
    for indice, empleado in enumerate(empleados):
        if cabecera_repetida and indice == 1:
            hoja.cell(row=fila, column=1, value="Reporte de Eventos de Asistencia")
            fila += 1

        hoja.cell(row=fila, column=1, value="ID:")
        hoja.cell(row=fila, column=3, value=empleado["id"])
        hoja.cell(row=fila, column=9, value="Nombre:")
        hoja.cell(row=fila, column=11, value=empleado["nombre"])
        hoja.cell(row=fila, column=19, value="Departamento:")
        hoja.cell(row=fila, column=21, value=empleado["departamento"])

        for dia, marcas in empleado["marcas"].items():
            hoja.cell(row=fila + 1, column=offset_dias + dia - 1, value=marcas)
        fila += 2

    libro.save(ruta)
    return ruta


@pytest.fixture(name="empleados")
def fixture_empleados():
    return [
        {
            "id": "1234567",
            "nombre": "JUAN PEREZ GOMEZ",
            "departamento": "MR LIN RESTAURANT",
            "marcas": {
                1: dt.time(8, 0),
                2: dt.time(8, 45),
                3: "08:00 12:00 13:00 18:00",
                # el dia 4 no tiene marcacion (ausencia)
                5: dt.time(7, 55),
            },
        },
        {
            "id": "7654321",
            "nombre": "MARIA LOPEZ BENITEZ",
            "departamento": "CAFETERIA CHICOLIN",
            "marcas": {1: dt.time(9, 10), 2: dt.time(8, 5)},
        },
    ]


@pytest.fixture(name="reporte")
def fixture_reporte(tmp_path, empleados):
    ruta = tmp_path / "planilla.xlsx"
    return construir_reporte(ruta, empleados)


# ---------------------------- Extraccion de horas ----------------------------


def test_extraer_horas_desde_time_datetime_texto_serial_y_multiple():
    assert extraer_horas(dt.time(8, 0)) == ["08:00"]
    assert extraer_horas(dt.datetime(1899, 12, 30, 18, 5)) == ["18:05"]
    assert extraer_horas("08:00") == ["08:00"]
    assert extraer_horas("08:00:30") == ["08:00"]
    assert extraer_horas("07:35 p.m.") == ["19:35"]
    assert extraer_horas(1 / 3) == ["08:00"]
    assert extraer_horas("08:00\n12:00\n13:00\n18:00") == ["08:00", "12:00", "13:00", "18:00"]
    assert extraer_horas("") == []
    assert extraer_horas(None) == []
    assert extraer_horas("FALTA") == []


# ---------------------------- Cabecera y periodo ----------------------------


@pytest.mark.parametrize("offset_dias", [1, 5])
def test_detecta_la_cabecera_de_dias_sin_importar_el_desplazamiento(tmp_path, empleados, offset_dias):
    ruta = construir_reporte(tmp_path / f"offset{offset_dias}.xlsx", empleados, offset_dias=offset_dias)
    grilla = leer_grilla(ruta, 0)

    fila, mapa = detectar_cabecera_dias(grilla)

    assert fila == 3  # fila 3 (0-based): la cabecera de dias del reporte
    assert sorted(mapa) == list(range(1, 32))
    assert mapa[1] == offset_dias - 1
    assert mapa[31] == offset_dias + 29


def test_detecta_el_periodo_y_lo_sobreescribe_con_mes(tmp_path, empleados):
    grilla = leer_grilla(construir_reporte(tmp_path / "periodo.xlsx", empleados), 0)

    periodo = detectar_periodo(grilla, 3)
    assert periodo == {"desde": "2026-08-01", "hasta": "2026-08-31", "anio": 2026, "mes": 8}

    _, sobreescrito, _, _ = parsear_reporte(grilla, (2027, 3))
    assert sobreescrito["anio"] == 2027
    assert sobreescrito["mes"] == 3


def test_usa_el_mes_indicado_cuando_el_archivo_no_trae_periodo(tmp_path, empleados):
    ruta = construir_reporte(tmp_path / "sinperiodo.xlsx", empleados, con_periodo=False)
    grilla = leer_grilla(ruta, 0)

    with pytest.raises(ValueError, match="periodo"):
        parsear_reporte(grilla)

    _, periodo, _, avisos = parsear_reporte(grilla, (2026, 8))
    assert periodo["anio"] == 2026 and periodo["mes"] == 8
    assert any("Periodo no detectado" in aviso for aviso in avisos)


# ---------------------------- Bloques por funcionario ----------------------------


def test_extrae_id_nombre_departamento_y_marcaciones_por_dia(reporte):
    grilla = leer_grilla(reporte, 0)
    funcionarios, periodo, mapa_dias, avisos = parsear_reporte(grilla)

    assert avisos == []
    assert len(funcionarios) == 2
    assert periodo["mes"] == 8 and len(mapa_dias) == 31

    juan = funcionarios[0]
    assert juan["id"] == "1234567"
    assert juan["nombre"] == "JUAN PEREZ GOMEZ"
    assert juan["departamento"] == "MR LIN RESTAURANT"
    assert juan["marcas"][1] == ["08:00"]
    assert juan["marcas"][3] == ["08:00", "12:00", "13:00", "18:00"]
    assert juan["marcas"][5] == ["07:55"]
    assert 4 not in juan["marcas"]  # dia sin marcacion

    maria = funcionarios[1]
    assert (maria["id"], maria["nombre"], maria["departamento"]) == (
        "7654321",
        "MARIA LOPEZ BENITEZ",
        "CAFETERIA CHICHOLIN".replace("CHICHOLIN", "CHICOLIN"),
    )
    assert maria["marcas"] == {1: ["09:10"], 2: ["08:05"]}


def test_ignora_encabezados_repetidos_por_paginacion(tmp_path, empleados):
    ruta = construir_reporte(tmp_path / "paginado.xlsx", empleados, cabecera_repetida=True)
    funcionarios, _, _, avisos = parsear_reporte(leer_grilla(ruta, 0))

    assert len(funcionarios) == 2
    assert any("paginacion" in aviso for aviso in avisos)


def test_fusiona_bloques_duplicados_del_mismo_id(tmp_path):
    empleados = [
        {"id": "999", "nombre": "ANA DIAZ", "departamento": "LIN GROUP", "marcas": {1: dt.time(8, 0)}},
        {"id": "999", "nombre": "ANA DIAZ", "departamento": "LIN GROUP", "marcas": {2: dt.time(8, 30)}},
    ]
    ruta = construir_reporte(tmp_path / "duplicado.xlsx", empleados)
    funcionarios, _, _, avisos = parsear_reporte(leer_grilla(ruta, 0))

    assert len(funcionarios) == 1
    assert funcionarios[0]["marcas"] == {1: ["08:00"], 2: ["08:30"]}
    assert any("repetido" in aviso for aviso in avisos)


def test_avisa_bloques_sin_fila_de_marcaciones(tmp_path, empleados):
    ruta = tmp_path / "sueltos.xlsx"
    libro = openpyxl.Workbook()
    hoja = libro.active
    for dia in range(1, 32):
        hoja.cell(row=4, column=4 + dia, value=dia)
    hoja.cell(row=5, column=1, value="ID:")
    hoja.cell(row=5, column=3, value="111")
    hoja.cell(row=5, column=9, value="Nombre:")
    hoja.cell(row=5, column=11, value="SIN MARCAS")
    hoja.cell(row=6, column=1, value="ID:")
    hoja.cell(row=6, column=3, value="222")
    hoja.cell(row=6, column=9, value="Nombre:")
    hoja.cell(row=6, column=11, value="OTRO")
    libro.save(ruta)

    funcionarios, _, _, avisos = parsear_reporte(leer_grilla(ruta, 0), (2026, 8))

    assert [f["id"] for f in funcionarios] == ["111", "222"]
    assert all(f["marcas"] == {} for f in funcionarios)
    assert sum("sin fila de marcaciones" in aviso for aviso in avisos) == 2


# ---------------------------- Salidas ----------------------------


def test_dataframe_largo_listo_para_el_importador_web(reporte):
    funcionarios, periodo, _, _ = parsear_reporte(leer_grilla(reporte, 0))
    largo = a_largo(funcionarios, periodo)

    assert list(largo.columns) == ["ID", "NOMBRE", "DEPARTAMENTO", "FECHA", "DIA", "ENTRADA", "SALIDA", "MARCADAS"]
    primera = largo.iloc[0].to_dict()
    assert (primera["ID"], primera["FECHA"], primera["ENTRADA"]) == ("1234567", "2026-08-01", "08:00")

    dia_tres = largo[(largo["ID"] == "1234567") & (largo["DIA"] == 3)].iloc[0]
    assert dia_tres["ENTRADA"] == "08:00"
    assert dia_tres["SALIDA"] == "18:00"
    assert dia_tres["MARCADAS"] == "08:00 12:00 13:00 18:00"

    # Juan marco 4 dias (1, 2, 3 y 5) y Maria 2 dias (1 y 2).
    assert len(largo) == 6


def test_incluir_vacios_agrega_los_dias_sin_marcacion(reporte):
    funcionarios, periodo, _, _ = parsear_reporte(leer_grilla(reporte, 0))

    assert len(a_largo(funcionarios, periodo, incluir_vacios=True)) == 62  # 2 funcionarios x 31 dias

    matriz = a_matriz(funcionarios, {dia: dia for dia in range(1, 32)}, periodo)
    assert matriz.shape == (2, 3 + 31)
    assert matriz.loc[0, "DIA_04"] == ""  # dia sin marcacion
    assert matriz.loc[1, "DIA_01"] == "09:10"


def test_json_expone_periodo_funcionarios_y_registros(reporte):
    funcionarios, periodo, _, _ = parsear_reporte(leer_grilla(reporte, 0))
    datos = a_json(funcionarios, periodo)

    assert datos["periodo"]["anio"] == 2026
    assert datos["funcionarios"][0]["diasConMarcacion"] == 4

    registro = next(item for item in datos["registros"] if item["fecha"] == "2026-08-03")
    assert registro == {
        "biometricId": "1234567",
        "nombre": "JUAN PEREZ GOMEZ",
        "fecha": "2026-08-03",
        "entrada": "08:00",
        "salida": "18:00",
        "marcas": ["08:00", "12:00", "13:00", "18:00"],
    }


# ---------------------------- CLI de punta a punta ----------------------------


def test_main_genera_los_archivos_de_salida(reporte, tmp_path, capsys):
    salida = tmp_path / "salida"

    assert main([str(reporte), "--out-dir", str(salida), "--json"]) == 0

    encabezado = (salida / "asistencia_largo.csv").read_text(encoding="utf-8-sig").splitlines()[0]
    assert encabezado == "ID;NOMBRE;DEPARTAMENTO;FECHA;DIA;ENTRADA;SALIDA;MARCADAS"
    assert (salida / "asistencia_matriz.csv").exists()

    datos = json.loads((salida / "asistencia.json").read_text(encoding="utf-8"))
    assert len(datos["registros"]) == 6

    resumen = capsys.readouterr().out
    assert "Periodo        : 2026-08-01 ~ 2026-08-31" in resumen
    assert "Funcionarios   : 2" in resumen
    assert "Avisos" not in resumen


def test_main_avisa_mes_invalido(reporte, tmp_path):
    assert main([str(reporte), "--out-dir", str(tmp_path / "x"), "--mes", "agosto"]) == 2
